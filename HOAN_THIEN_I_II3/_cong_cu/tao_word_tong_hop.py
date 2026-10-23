"""Ghép báo cáo Phần I và Phần II đến hết II.3 từ các bài đã sửa.

- Phần I: bản Phần I đã sửa (giữ nguyên định dạng gốc, dùng làm khuôn style).
- II.1: UCD tổng quát, phân cấp chức năng, bảng chức năng con (catalogue.json).
- II.2: mỗi vùng có UCD vùng; UC trọng tâm lấy nguyên mục 2–5 của bài thành viên đã sửa
  (bỏ mục 1 "Thông tin phân công", bỏ II.2+ trở đi); UC còn lại là FDS tóm tắt (fds.json);
  V10 lấy từ v10.json kèm AD draw.io.
- II.3: tổng quan gói, năm hình cụm lớp, bảng danh mục lớp.

    python3 tao_word_tong_hop.py <thư mục HOAN_THIEN_I_II3> <file .docx ra>
"""
import copy
import importlib.util
import json
import os
import re
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docxcompose.composer import Composer

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
DATA = os.path.join(REPO, "PHAN_II_RA_SOAT", "_cong_cu", "du_lieu")


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


W = load("w", os.path.join(REPO, "PHAN_II_RA_SOAT", "_cong_cu", "tao_bao_cao_word.py"))
G3 = load("g3", os.path.join(HERE, "gen_ii3_drawio.py"))
W.MEANING["YeuCauNguonLuc"] = "Yêu cầu nguồn lực PM gửi cho V5 kèm khung thời gian (CN-127)"

P1 = "[PTTKHT] I - PHÂN TÍCH YÊU CẦU FINAL (đã sửa).docx"
# UC trọng tâm -> (bài đã sửa, tiêu đề mục); bài V2b gồm cả 01.2 và 01.3.
MEMBER = {
    "UC-01.1": ("V2-Phần II (đã sửa).docx", "V2a – Nguyễn Thị Thùy Linh"),
    "UC-01.2": ("V2-Phần II (2) - ban 22-10 (đã sửa).docx", "V2b – Nguyễn Tuấn Long"),
    "UC-02.1": ("V3- Phần 2 (đã sửa).docx", "V3 – Lê Hoài Linh"),
    "UC-03.3": ("V4 - Phần II (đã sửa).docx", "V4 – Nguyễn Hoàng Linh"),
    "UC-04.1": ("NV05 I2 - I3 (đã sửa).docx", "V5 – Nguyễn Hoàng Khánh Linh"),
    "UC-05.1": ("V6 - Phần II (đã sửa).docx", "V6 – Nguyễn Thi Thùy Linh"),
    "UC-06.2": ("MNV7_II_Thiết kế Logic (đã sửa).docx", "V7 – Lưu Thùy Linh"),
    "UC-07.1": ("MNV8 - Ngọc Linh (đã sửa).docx", "V8 – Ngọc Linh"),
    "UC-08.1": ("V9 - Phần II (đã sửa).docx", "V9 – Ngô Nhật Long"),
}
COVERED = {"UC-01.3": "UC-01.2"}


def text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def blank(template):
    """Tài liệu rỗng mang style của Phần I."""
    d = Document(template)
    body = d.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    return d


def widen(tbl, full=9300):
    """Bảng của bài thành viên hẹp hơn nhiều so với khổ trang thì giãn cột theo tỷ lệ cho đủ khổ."""
    cols = tbl.findall(qn("w:tblGrid") + "/" + qn("w:gridCol"))
    ws = [float(c.get(qn("w:w"), "0")) for c in cols]
    total = sum(ws)
    if not total or total >= 7000:
        return
    k = full / total
    for c, w in zip(cols, ws):
        c.set(qn("w:w"), str(int(w * k)))
    for tcw in tbl.iter(qn("w:tcW")):
        if tcw.get(qn("w:type")) == "dxa":
            tcw.set(qn("w:w"), str(int(float(tcw.get(qn("w:w"), "0")) * k)))
    tblw = tbl.find(qn("w:tblPr") + "/" + qn("w:tblW"))
    if tblw is not None:
        tblw.set(qn("w:type"), "dxa")
        tblw.set(qn("w:w"), str(full))


def extract(path, fig_ref, tmp):
    """Mục 2–5 của bài thành viên. Tiêu đề thành đoạn in đậm; hình UCD thay bằng dẫn chiếu hình vùng."""
    d = Document(path)
    body = d.element.body
    els = [e for e in body if e.tag != qn("w:sectPr")]
    start = next(i for i, e in enumerate(els) if text(e).startswith("2.") and "Use Case Diagram" in text(e))
    end = next(i for i, e in enumerate(els) if text(e).upper().startswith("II.2+"))
    in_ucd, noted = False, False
    for i, e in enumerate(els):
        if i < start or i >= end:
            body.remove(e)
            continue
        if e.tag != qn("w:p"):
            continue
        t = text(e)
        if t.startswith("2.3"):
            in_ucd = True
        elif t.startswith("3."):
            in_ucd = False
        if in_ucd and list(e.iter(qn("w:drawing"))):
            if noted:
                body.remove(e)
                continue
            for r in list(e):
                if r.tag in (qn("w:r"), qn("w:hyperlink")):
                    e.remove(r)
            r = e.makeelement(qn("w:r"), {})
            rpr = r.makeelement(qn("w:rPr"), {})
            rpr.append(rpr.makeelement(qn("w:i"), {}))
            r.append(rpr)
            tt = r.makeelement(qn("w:t"), {})
            tt.text = f"Sơ đồ ca sử dụng của vùng: xem {fig_ref}."
            r.append(tt)
            e.append(r)
            noted = True
            continue
        ppr = e.find(qn("w:pPr"))
        sty = ppr.find(qn("w:pStyle")) if ppr is not None else None
        heading = sty is not None and sty.get(qn("w:val"), "").lower().startswith(("heading", "title"))
        numbered = bool(re.match(r"^[2-5](\.\d+)*\.\s", t)) and len(t) < 90
        if heading or numbered:
            if sty is not None:
                ppr.remove(sty)
            num = ppr.find(qn("w:numPr")) if ppr is not None else None
            if num is not None:
                ppr.remove(num)
            if ppr is None:
                ppr = e.makeelement(qn("w:pPr"), {})
                e.insert(0, ppr)
            ppr.append(ppr.makeelement(qn("w:keepNext"), {}))
            for r in e.iter(qn("w:r")):
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = r.makeelement(qn("w:rPr"), {})
                    r.insert(0, rpr)
                if rpr.find(qn("w:b")) is None:
                    rpr.append(rpr.makeelement(qn("w:b"), {}))
    for tbl in body.iter(qn("w:tbl")):
        widen(tbl)
    out = os.path.join(tmp, os.path.basename(path))
    d.save(out)
    return out


def build(root, out):
    root = os.path.abspath(root)
    src = os.path.join(root, "bai_goc_da_sua")
    p1 = os.path.join(src, P1)
    dr = os.path.join(root, "drawio")
    cat = json.load(open(os.path.join(DATA, "catalogue.json"), encoding="utf-8"))
    fds = json.load(open(os.path.join(DATA, "fds.json"), encoding="utf-8"))
    v10 = json.load(open(os.path.join(DATA, "v10.json"), encoding="utf-8"))
    tmp = tempfile.mkdtemp()
    v10root = os.path.join(tmp, "v10")
    os.makedirs(os.path.join(v10root, "drawio"))
    os.symlink(os.path.join(dr, "II.2_AD"), os.path.join(v10root, "drawio", "II.2_V10"))
    parts = []          # (loại, đường dẫn) theo thứ tự ghép
    cur = [None]

    def chunk():
        if cur[0] is None:
            cur[0] = blank(p1)
        return cur[0]

    def flush():
        if cur[0] is not None:
            path = os.path.join(tmp, f"chunk_{len(parts):02d}.docx")
            cur[0].save(path)
            parts.append(path)
            cur[0] = None

    doc = chunk()
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("II — PHÂN TÍCH (THIẾT KẾ LOGIC)", level=1)

    # ---------------- II.1
    doc.add_heading("II.1 — PHẠM VI HỆ THỐNG CÔNG NGHỆ", level=2)
    W.para(doc, "Hệ thống quản lý và điều phối nguồn lực dự án của Lạc Việt Consulting hỗ trợ toàn bộ quá trình "
                "từ xác định nhu cầu nguồn lực đến theo dõi và dự báo cung – cầu. Phạm vi gồm 9 chức năng cấp vùng "
                f"(UC-01 đến UC-09), tương ứng các vùng nghiệp vụ V2–V10 của Phần I, và {len(cat['ucs'])} chức "
                "năng con. Tác nhân gồm PM, RM, Nguồn lực, Tài chính, PMO/Quản lý danh mục và CEO/Managing "
                "Director.")
    W.figure(doc, os.path.join(dr, "II.1", "II1_UC_Tong_Quat.png"), "Hình II.1-1. Sơ đồ ca sử dụng tổng quát.",
             width=14)
    W.figure(doc, os.path.join(dr, "II.1", "Phan_cap_chuc_nang.png"), "Hình II.1-2. Sơ đồ phân cấp chức năng.",
             width=16)
    W.para(doc, f"Bảng II.1-1 liệt kê {len(cat['ucs'])} chức năng con. 12 ca sử dụng trọng tâm được đặc tả đầy đủ "
                "kèm sơ đồ hoạt động; các ca sử dụng còn lại được mô tả bằng FDS tóm tắt.")
    W.caption(doc, "Bảng II.1-1. Danh sách chức năng con.", before=True)
    W.table(doc, ["Vùng", "Mã", "Chức năng", "Tác nhân", "Chức năng Phần I", "Mức đặc tả"],
            [(u["region"], u["id"], u["name"], u["actor"].replace(";", ","), "CN-" + u["cn"],
              "Đặc tả đầy đủ, AD" if u["id"] in W.FOCUS else "FDS tóm tắt") for u in cat["ucs"]],
            widths=[1.3, 1.9, 5, 2.6, 3, 2.2], size=10)

    # ---------------- II.2
    doc.add_heading("II.2 — MÔ HÌNH HÓA CHỨC NĂNG CỦA HỆ THỐNG CÔNG NGHỆ", level=2)
    W.para(doc, "Mỗi vùng gồm sơ đồ ca sử dụng của vùng; đặc tả đầy đủ (Use Case Description) và sơ đồ hoạt động "
                "cho ca sử dụng trọng tâm, lấy từ bài của thành viên phụ trách sau khi đã sửa; FDS tóm tắt cho các "
                "ca sử dụng còn lại.")
    fig = 0
    for k, (code, vname, v, subs) in enumerate(cat["regions"], 1):
        doc = chunk()
        doc.add_heading(f"II.2.{k}. {v} — {vname} ({code})", level=3)
        fig += 1
        ucd_ref = f"Hình II.2-{fig}"
        W.figure(doc, os.path.join(dr, "II.2_UCD_vung", f"UCD_{code}.png"),
                 f"{ucd_ref}. Sơ đồ ca sử dụng {code} {vname}.", width=15)
        for cid, cname in subs:
            if cid in COVERED:
                continue
            if cid in MEMBER:
                fname, who = MEMBER[cid]
                title = f"{cid} {cname}"
                if cid == "UC-01.2":
                    title = "UC-01.2 Xác nhận nhu cầu và UC-01.3 Điều chỉnh nhu cầu theo phản hồi"
                doc.add_heading(f"{title} (ca sử dụng trọng tâm)", level=4)
                W.para(doc, f"Người thực hiện: {who}.", italic=True)
                flush()
                parts.append(extract(os.path.join(src, fname), ucd_ref, tmp))
            elif v == "V10":
                doc = chunk()
                doc.add_heading(f"{cid} {cname} (ca sử dụng trọng tâm)", level=4)
                s = next(x for x in v10["specs"] if x["id"] == cid)
                fig += 1
                W.v10_spec(doc, s, fig, v10root)
            else:
                doc = chunk()
                doc.add_heading(f"{cid} {cname} (FDS tóm tắt)", level=4)
                rows = W.clean_fds(cid, fds[cid])
                if cid == "UC-08.2":
                    rows = [(a, b.replace("theo công thức chung.", "theo công thức chung [Chờ xác nhận QĐ-07].")
                             if a == "Luồng chính" else b) for a, b in rows]
                W.kv_table(doc, rows)
        if v == "V10":
            doc = chunk()
            W.para(doc, "Quy tắc nghiệp vụ áp dụng cho UC-09.1 và UC-09.2", bold=True)
            W.table(doc, ["Căn cứ", "Nội dung"], [tuple(r) for r in v10["rules"]], widths=[4.2, 11.8])

    # ---------------- II.3
    doc = chunk()
    doc.add_heading("II.3 — MÔ HÌNH HÓA CẤU TRÚC TĨNH CỦA HỆ THỐNG CÔNG NGHỆ", level=2)
    n = sum(len(c[4]) for c in G3.CLUSTERS)
    W.para(doc, f"Sơ đồ lớp miền nghiệp vụ gồm {n} lớp, chia thành năm gói theo nhóm vùng. Hình II.3-1 cho thấy các "
                "gói và quan hệ phụ thuộc giữa chúng. Hình II.3-2 đến II.3-6 trình bày từng gói với thuộc tính, liên "
                "kết và bội số. Lớp thuộc gói khác được vẽ nền xám, viền nét đứt, chỉ để thể hiện liên kết; thuộc "
                "tính của lớp đó xem ở hình của gói sở hữu. Bảng lớp, thuộc tính và bội số trong phần đề xuất lớp "
                "(II.2+) của từng thành viên đã được sửa cho khớp các hình này.")
    files = ["II3-1_Tong_quan_goi_lop", "II3-2_Nhu_cau_va_nang_luc", "II3-3_Tim_kiem_kha_nang_tham_gia_lich",
             "II3-4_Phuong_an_xung_dot_phe_duyet", "II3-5_Chi_phi_ngan_sach", "II3-6_Theo_doi_du_bao"]
    W.figure(doc, os.path.join(dr, "II.3", files[0] + ".png"), "Hình II.3-1. Tổng quan các gói lớp miền nghiệp vụ.",
             width=15)
    for k, ((key, fno, title, vung, _), f) in enumerate(zip(G3.CLUSTERS, files[1:]), 1):
        doc.add_heading(f"II.3.{k}. Gói {key}: {title} ({vung})", level=3)
        W.figure(doc, os.path.join(dr, "II.3", f + ".png"), f"Hình {fno}. Sơ đồ lớp miền: {title.lower()}.",
                 width=16)
    doc.add_heading("II.3.6. Danh mục lớp", level=3)
    W.caption(doc, "Bảng II.3-1. Danh mục lớp miền nghiệp vụ.", before=True)
    rows = []
    for key, fno, _, vung, cs in G3.CLUSTERS:
        for c in cs:
            rows.append((c, W.MEANING[c], ", ".join(G3.CLASSES[c]), fno))
    W.table(doc, ["Lớp", "Ý nghĩa", "Thuộc tính", "Hình"], rows, widths=[4.4, 4.6, 5.4, 1.6], size=10)
    flush()

    base = Document(p1)
    comp = Composer(base)
    for p in parts:
        comp.append(Document(p))
    comp.save(out)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
