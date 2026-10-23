"""Tiện ích Word và danh mục dùng chung cho tao_word_tong_hop.py: bảng, chú thích, hình, ý nghĩa lớp II.3,
danh sách UC trọng tâm và người phụ trách."""
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FOCUS = ["UC-01.1", "UC-01.2", "UC-01.3", "UC-02.1", "UC-03.3", "UC-04.1", "UC-05.1", "UC-06.2",
         "UC-07.1", "UC-08.1", "UC-09.1", "UC-09.2"]
# Người phụ trách UC trọng tâm (HD-PCCN §5); V10 do trưởng nhóm phụ trách (QĐ-23).
OWNER = {"UC-01.1": "V2a – Nguyễn Thị Thùy Linh", "UC-01.2": "V2b – Nguyễn Tuấn Long",
         "UC-01.3": "V2b – Nguyễn Tuấn Long", "UC-02.1": "V3 – Lê Hoài Linh",
         "UC-03.3": "V4 – Nguyễn Hoàng Linh", "UC-04.1": "V5 – Nguyễn Hoàng Khánh Linh",
         "UC-05.1": "V6 – Nguyễn Thi Thùy Linh", "UC-06.2": "V7 – Lưu Thùy Linh",
         "UC-07.1": "V8 – Ngọc Linh", "UC-08.1": "V9 – Ngô Nhật Long",
         "UC-09.1": "V10 – LUẤN", "UC-09.2": "V10 – LUẤN"}

# Ý nghĩa lớp: theo danh mục II.3 của bản tích hợp; lớp V9, V10 mô tả theo thuộc tính.
MEANING = {
    "DuAn": "Dự án dùng nguồn lực, có hạng ưu tiên",
    "GiaiDoanDuAn": "Giai đoạn của dự án",
    "NhuCauNguonLuc": "Nhu cầu nguồn lực do PM lập, RM xác nhận",
    "ViTriCanBoTri": "Vị trí trong nhu cầu: vai trò, số lượng, %FTE",
    "YeuCauKyNang": "Kỹ năng mà vị trí đòi hỏi",
    "PhanHoiNhuCau": "Phản hồi từ V4, V10 để PM xem lại nhu cầu",
    "NguonLuc": "Người có thể bố trí: nội bộ, cộng tác viên/đối tác, tuyển ngoài",
    "KyNang": "Danh mục kỹ năng",
    "HoSoNangLuc": "Hồ sơ năng lực tập trung của một nguồn lực",
    "NangLucKyNang": "Mức thành thạo một kỹ năng và trạng thái xác nhận",
    "BoTieuChiNangLuc": "Bộ tiêu chí đánh giá năng lực chung",
    "TieuChiNangLuc": "Một tiêu chí năng lực",
    "MucDoNangLuc": "Mức của tiêu chí và điều kiện công nhận",
    "BoTieuChiPhuHop": "Bộ điều kiện bắt buộc, tiêu chí, trọng số cho một nhu cầu",
    "TieuChiPhuHop": "Điều kiện bắt buộc hoặc tiêu chí có trọng số",
    "PhienTimKiem": "Một lần tìm kiếm và đánh giá cho một nhu cầu",
    "KetQuaDanhGiaPhuHop": "Phân loại, điểm, thứ hạng, vai trò shortlist của một ứng viên",
    "DiemTieuChi": "Điểm theo từng tiêu chí",
    "DeNghiMienTru": "Đề nghị miễn trừ điều kiện bắt buộc và kết quả duyệt",
    "LichNguonLuc": "Lịch nguồn lực tập trung của một nguồn lực",
    "PhanBoNguonLuc": "Bản ghi phân bổ cứng hoặc mềm trên lịch",
    "KetLuanKhaNangThamGia": "Kết luận khả năng tham gia, mức sử dụng, cờ rủi ro",
    "DeXuatThamGiaMotPhan": "Tỷ lệ và thời gian đề xuất khi chỉ tham gia được một phần",
    "HoSoVuotTaiNgoaiLe": "Trường hợp vượt tải chuyển cấp theo DoA",
    "PhuongAnBoTri": "Phương án chính hoặc thay thế, có phiên bản và trạng thái",
    "ChiTietBoTri": "Một dòng bố trí: người, vị trí, giai đoạn, tỷ lệ",
    "RuiRoViTri": "Rủi ro và biện pháp của vị trí rủi ro cao",
    "DiemXungDot": "Điểm vượt ngưỡng công suất, có mức cảnh báo",
    "NhatKyXuLyXungDot": "Nhật ký xử lý điểm xung đột, đàm phán, phân xử",
    "PhanXuXungDot": "Hồ sơ chuyển cấp và quyết định phân xử",
    "DeXuatThayDoi": "Đề xuất thay đổi tiến độ hoặc phạm vi do cân đối nguồn lực",
    "HoSoTrinhDuyet": "Một lần trình duyệt phương án: mức rủi ro, cấp thẩm quyền",
    "YeuCauGiaiTrinh": "Yêu cầu RM giải trình, hoàn thiện phương án",
    "QuyetDinhPheDuyet": "Quyết định, lý do, điều kiện, người và thời điểm quyết định",
    "DonGia": "Đơn giá áp dụng theo loại nguồn lực",
    "DuToanChiPhi": "Dự toán chi phí của phương án",
    "DuBaoChiPhi": "Lần cập nhật dự báo chi phí",
    "ChiPhiThucTe": "Chi phí thực tế phát sinh",
    "NganSach": "Ngân sách được duyệt và phần còn lại",
    "DanhGiaChiPhi": "Đối chiếu chi phí với ngân sách tại một thời điểm",
    "NguyenNhanVuot": "Nguyên nhân vượt ngân sách",
    "HuongXuLy": "Hướng xử lý khi vượt ngân sách",
    "SoSanhPhuongAn": "Một lần so sánh các phương án",
    "ChiTietSoSanh": "Số liệu của một phương án trong lần so sánh",
    "KetQuaDuBao": "Kết quả một lần tổng hợp và dự báo cung – cầu",
    "TongHopNguonCung": "Nguồn cung có thể đáp ứng trong khoảng thời gian",
    "NhuCauDuBao": "Nhu cầu được xét đưa vào dự báo",
    "DongDoiChieuCungCau": "Đối chiếu cung – cầu theo từng khoảng thời gian",
}


def shade(cell, hexcolor):
    tc = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hexcolor)
    tc.append(s)


def para(doc, text, bold=False, italic=False, align=None, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    if bold:
        p.paragraph_format.keep_with_next = True
    return p


def table(doc, header, rows, widths=None, size=11):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    trpr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = str(v).split("\n")
            for j, line in enumerate(lines):
                p = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                r = p.add_run(line)
                r.font.size = Pt(size)
    widths = widths or [16 / len(header)] * len(header)
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def caption(doc, text, before=False):
    p = para(doc, text, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    p.paragraph_format.keep_with_next = before
    p.paragraph_format.space_after = Pt(12)


def figure(doc, path, cap, width=16):
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, cap)


def kv_table(doc, rows):
    table(doc, ["Thành phần", "Nội dung"], rows, widths=[4.2, 11.8])
