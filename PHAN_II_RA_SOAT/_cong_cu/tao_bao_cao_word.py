"""Dựng bản Word Phần II (bản rà soát) từ hình draw.io đã xuất và dữ liệu trong du_lieu/.

- II.1: UCD tổng quát, phân cấp chức năng, bảng chức năng con (catalogue.json).
- II.2: UCD vùng; FDS tóm tắt cho UC không trọng tâm (fds.json, đã bỏ phần chưa có nguồn);
  UC trọng tâm để khung [Cần bổ sung] cho thành viên, trừ V4 (giữ đặc tả của bài V4) và
  V10 (v10.json, kèm AD draw.io).
- II.3: tổng quan gói và năm hình cụm lớp, bảng danh mục lớp.
- II.4: danh sách sơ đồ đã chọn theo II4_RA_SOAT_VA_PHAN_CONG.md, mỗi sơ đồ một khung chờ.

    python3 tao_bao_cao_word.py <thư mục PHAN_II_RA_SOAT> <file .docx ra>
"""
import importlib.util
import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "du_lieu")

FOCUS = ["UC-01.1", "UC-01.2", "UC-01.3", "UC-02.1", "UC-03.3", "UC-04.1", "UC-05.1", "UC-06.2",
         "UC-07.1", "UC-08.1", "UC-09.1", "UC-09.2"]
# Người phụ trách UC trọng tâm (HD-PCCN §5; bảng II.2 của file yêu cầu sửa).
OWNER = {"UC-01.1": "V2a – Nguyễn Thị Thùy Linh", "UC-01.2": "V2b – Nguyễn Tuấn Long",
         "UC-01.3": "V2b – Nguyễn Tuấn Long", "UC-02.1": "V3 – Lê Hoài Linh",
         "UC-03.3": "V4 – Nguyễn Hoàng Linh", "UC-04.1": "V5 – Nguyễn Hoàng Khánh Linh",
         "UC-05.1": "V6 – Nguyễn Thi Thùy Linh", "UC-06.2": "V7 – Lưu Thùy Linh",
         "UC-07.1": "V8 – Ngọc Linh", "UC-08.1": "V9 – Ngô Nhật Long"}
FIX_TABLE = {"UC-01.1": "V2a", "UC-01.2": "V2b", "UC-01.3": "V2b", "UC-02.1": "V3", "UC-03.3": "V4",
             "UC-04.1": "V5", "UC-05.1": "V6", "UC-06.2": "V7", "UC-07.1": "V8", "UC-08.1": "V9"}

# Ý nghĩa lớp: theo danh mục II.3 của bản tích hợp; lớp V9, V10 mô tả theo thuộc tính.
MEANING = {
    "DuAn": "Dự án dùng nguồn lực, có hạng ưu tiên",
    "GiaiDoanDuAn": "Giai đoạn của dự án",
    "NhuCauNguonLuc": "Nhu cầu nguồn lực do PM lập, RM xác nhận",
    "ViTriCanBoTri": "Vị trí trong nhu cầu: vai trò, số lượng, %FTE",
    "YeuCauKyNang": "Kỹ năng mà vị trí đòi hỏi",
    "PhanHoiNhuCau": "Phản hồi từ V4, V10 để PM xem lại nhu cầu",
    "NguonLuc": "Người có thể bố trí: nội bộ, cộng tác viên/đối tác, tuyển ngoài",
    "KyNang": "Danh mục kỹ năng",
    "HoSoNangLuc": "Hồ sơ năng lực tập trung của một nguồn lực",
    "NangLucKyNang": "Mức thành thạo một kỹ năng và trạng thái xác nhận",
    "BoTieuChiNangLuc": "Bộ tiêu chí đánh giá năng lực chung",
    "TieuChiNangLuc": "Một tiêu chí năng lực",
    "MucDoNangLuc": "Mức của tiêu chí và điều kiện công nhận",
    "BoTieuChiPhuHop": "Bộ điều kiện bắt buộc, tiêu chí, trọng số cho một nhu cầu",
    "TieuChiPhuHop": "Điều kiện bắt buộc hoặc tiêu chí có trọng số",
    "PhienTimKiem": "Một lần tìm kiếm và đánh giá cho một nhu cầu",
    "KetQuaDanhGiaPhuHop": "Phân loại, điểm, thứ hạng, vai trò shortlist của một ứng viên",
    "DiemTieuChi": "Điểm theo từng tiêu chí",
    "DeNghiMienTru": "Đề nghị miễn trừ điều kiện bắt buộc và kết quả duyệt",
    "LichNguonLuc": "Lịch nguồn lực tập trung của một nguồn lực",
    "PhanBoNguonLuc": "Bản ghi phân bổ cứng hoặc mềm trên lịch",
    "KetLuanKhaNangThamGia": "Kết luận khả năng tham gia, mức sử dụng, cờ rủi ro",
    "DeXuatThamGiaMotPhan": "Tỷ lệ và thời gian đề xuất khi chỉ tham gia được một phần",
    "HoSoVuotTaiNgoaiLe": "Trường hợp vượt tải chuyển cấp theo DoA",
    "PhuongAnBoTri": "Phương án chính hoặc thay thế, có phiên bản và trạng thái",
    "ChiTietBoTri": "Một dòng bố trí: người, vị trí, giai đoạn, tỷ lệ",
    "RuiRoViTri": "Rủi ro và biện pháp của vị trí rủi ro cao",
    "DiemXungDot": "Điểm vượt ngưỡng công suất, có mức cảnh báo",
    "NhatKyXuLyXungDot": "Nhật ký xử lý điểm xung đột, đàm phán, phân xử",
    "PhanXuXungDot": "Hồ sơ chuyển cấp và quyết định phân xử",
    "DeXuatThayDoi": "Đề xuất thay đổi tiến độ hoặc phạm vi do cân đối nguồn lực",
    "HoSoTrinhDuyet": "Một lần trình duyệt phương án: mức rủi ro, cấp thẩm quyền",
    "YeuCauGiaiTrinh": "Yêu cầu RM giải trình, hoàn thiện phương án",
    "QuyetDinhPheDuyet": "Quyết định, lý do, điều kiện, người và thời điểm quyết định",
    "DonGia": "Đơn giá áp dụng theo loại nguồn lực",
    "DuToanChiPhi": "Dự toán chi phí của phương án",
    "DuBaoChiPhi": "Lần cập nhật dự báo chi phí",
    "ChiPhiThucTe": "Chi phí thực tế phát sinh",
    "NganSach": "Ngân sách được duyệt và phần còn lại",
    "DanhGiaChiPhi": "Đối chiếu chi phí với ngân sách tại một thời điểm",
    "NguyenNhanVuot": "Nguyên nhân vượt ngân sách",
    "HuongXuLy": "Hướng xử lý khi vượt ngân sách",
    "SoSanhPhuongAn": "Một lần so sánh các phương án",
    "ChiTietSoSanh": "Số liệu của một phương án trong lần so sánh",
    "KetQuaDuBao": "Kết quả một lần tổng hợp và dự báo cung – cầu",
    "TongHopNguonCung": "Nguồn cung có thể đáp ứng trong khoảng thời gian",
    "NhuCauDuBao": "Nhu cầu được xét đưa vào dự báo",
    "DongDoiChieuCungCau": "Đối chiếu cung – cầu theo từng khoảng thời gian",
}

# Sơ đồ II.4 đã chọn (II4_RA_SOAT_VA_PHAN_CONG.md, bỏ các dòng Bỏ và Chờ quyết định).
SM = [("NhuCauNguonLuc", "UC-01.1–01.3", "V2b – Nguyễn Tuấn Long", "CN-03, CN-11–CN-14, CN-17",
       "SM_NhuCauNguonLuc"),
      ("PhuongAnBoTri", "UC-05.1–05.2, UC-07.3", "V6 – Nguyễn Thi Thùy Linh (V8 góp phần sau trình)",
       "CN-60–CN-73, CN-99; QT-31", "SM_PhuongAnBoTri"),
      ("DiemXungDot", "UC-06.1–06.3", "V7 – Lưu Thùy Linh", "CN-74–CN-88; QT-33–QT-39",
       "SM_DiemXungDot"),
      ("HoSoTrinhDuyet", "UC-07.1–07.4", "V8 – Ngọc Linh", "CN-89–CN-102; QT-40–QT-46",
       "SM_HoSoTrinhDuyet")]
COM = [("Tra cứu căn cứ đánh giá ứng viên", "UC-03.3", "V4 – Nguyễn Hoàng Linh",
        "CN-34–CN-42; QT-15, QT-16", "COM_UC-03.3"),
       ("Tra cứu số liệu so sánh phương án", "UC-08.3", "V9 – Ngô Nhật Long", "CN-110; QT-50",
        "COM_UC-08.3"),
       ("Tra cứu tình trạng cung – cầu", "UC-09.1", "Trưởng nhóm", "CN-123; PCN-08", "COM_UC-09.1")]
SEQ = [("Ghi nhu cầu nguồn lực", "UC-01.1", "V2a – Nguyễn Thị Thùy Linh", "CN-01–CN-09, CN-11"),
       ("Xác nhận và bàn giao nhu cầu", "UC-01.2", "V2b – Nguyễn Tuấn Long", "CN-12–CN-14, CN-16"),
       ("Đề nghị và duyệt tiêu chí tạm", "UC-02.1", "V3 – Lê Hoài Linh", "CN-18, CN-19"),
       ("Xét miễn trừ điều kiện bắt buộc", "UC-03.3", "V4 – Nguyễn Hoàng Linh",
        "CN-34–CN-37; QT-13, QT-14"),
       ("Ghi kết luận khả năng tham gia", "UC-04.1", "V5 – Nguyễn Hoàng Khánh Linh",
        "CN-47–CN-52, CN-59, CN-127"),
       ("Lập và gửi phương án cho PM", "UC-05.1", "V6 – Nguyễn Thi Thùy Linh", "CN-60–CN-68, CN-73"),
       ("Cân đối, leo thang và bàn giao", "UC-06.2", "V7 – Lưu Thùy Linh",
        "CN-78–CN-82, CN-86–CN-88"),
       ("Rà soát hồ sơ trước quyết định", "UC-07.1", "V8 – Ngọc Linh", "CN-89–CN-94"),
       ("Tính dự toán chi phí phương án", "UC-08.1", "V9 – Ngô Nhật Long", "CN-103–CN-107"),
       ("Lưu kết quả dự báo theo thời gian", "UC-09.2", "Trưởng nhóm",
        "CN-116–CN-119, CN-125, CN-126")]

PH = "[Cần bổ sung]"


# ---------------------------------------------------------------- tiện ích Word
def shade(cell, hexcolor):
    tc = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hexcolor)
    tc.append(s)


def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(13)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.3
    for i, size in [(1, 16), (2, 14), (3, 13), (4, 13)]:
        h = doc.styles[f"Heading {i}"]
        h.font.name = "Times New Roman"
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = i == 4
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)


def para(doc, text, bold=False, italic=False, align=None, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    if bold:
        p.paragraph_format.keep_with_next = True
    return p


def placeholder(doc, who, what):
    p = doc.add_paragraph()
    r = p.add_run(f"[Cần bổ sung — {who}] ")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(what).italic = True
    return p


def table(doc, header, rows, widths=None, size=11):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    trpr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = str(v).split("\n")
            for j, line in enumerate(lines):
                p = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                red = line.startswith("[Cần") or line.startswith("[Chờ")
                r = p.add_run(line)
                r.font.size = Pt(size)
                if red:
                    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    widths = widths or [16 / len(header)] * len(header)
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def caption(doc, text, before=False):
    p = para(doc, text, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    p.paragraph_format.keep_with_next = before
    p.paragraph_format.space_after = Pt(12)


def figure(doc, path, cap, width=16):
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, cap)


def kv_table(doc, rows):
    table(doc, ["Thành phần", "Nội dung"], rows, widths=[4.2, 11.8])


# ---------------------------------------------------------------- nội dung
def clean_fds(uc, rows):
    out = []
    for k, v in rows[1:]:
        if uc == "UC-08.2" and k.startswith("Ràng buộc"):
            v = ("QT-47; CN-108–109: đối chiếu chi phí với ngân sách duyệt và phần còn lại.\n"
                 "[Chờ chốt QĐ-07: cách tính tỷ lệ vượt ngân sách.]\nTruy vết: I FINAL CN-108–109")
        out.append((k, v))
    return out


def focus_template(doc, uc):
    cid, name = uc["id"], uc["name"]
    placeholder(doc, OWNER[cid],
                f"Chèn đặc tả và sơ đồ hoạt động đã sửa theo bảng {FIX_TABLE[cid]} của file "
                "II2_YEU_CAU_SUA_THEO_VUNG. Chỉ chèn khi mọi dòng của bảng đạt tiêu chí.")
    kv_table(doc, [
        ("Tên ca sử dụng", name), ("ID", cid), ("Tác nhân liên quan", uc["actor"].replace(";", ",")),
        ("Phạm vi (theo bảng II.1)", uc["scope"]),
        ("Mục tiêu", PH), ("Sự kiện kích hoạt", PH), ("Tiền điều kiện", PH), ("Hậu điều kiện", PH),
        ("Luồng chính", PH), ("Luồng thay thế", PH), ("Luồng ngoại lệ", PH),
        ("Quy tắc nghiệp vụ", PH), ("Truy vết", "I FINAL CN-" + uc["cn"]),
    ])
    placeholder(doc, OWNER[cid], f"Sơ đồ hoạt động {cid} {name}: chèn PNG xuất từ AD_{cid}.drawio.")


def v4_33(doc):
    para(doc, "Đặc tả dưới đây lấy từ bài V4. Các ô đánh dấu đỏ còn phải sửa theo bảng V4 của "
              "file II2_YEU_CAU_SUA_THEO_VUNG.", italic=True)
    kv_table(doc, [
        ("Tên ca sử dụng", "Đánh giá, xếp hạng và lập shortlist ứng viên"), ("ID", "UC-03.3"),
        ("Mục tiêu", "RM đánh giá mức độ phù hợp của các ứng viên đã tìm được theo điều kiện bắt buộc, "
                     "tiêu chí và trọng số của nhu cầu, xếp hạng họ và lập shortlist gồm ứng viên chính "
                     "và ứng viên thay thế."),
        ("Tác nhân chính", "RM"),
        ("Tác nhân phụ", "PMO/Quản lý danh mục (phê duyệt miễn trừ điều kiện bắt buộc)"),
        ("Sự kiện kích hoạt", "RM bắt đầu đánh giá một danh sách ứng viên đã có từ UC-03.2 cho một "
                              "nhu cầu nguồn lực."),
        ("Tiền điều kiện", "(1) Nhu cầu ở trạng thái “đã xác nhận” (QT-12).\n"
                           "(2) Đã có bộ điều kiện bắt buộc, tiêu chí và trọng số cho nhu cầu (UC-03.1, QT-19).\n"
                           "(3) Đã có danh sách ứng viên (UC-03.2).\n"
                           "(4) Hồ sơ năng lực tập trung của V3 tra cứu được."),
        ("Hậu điều kiện (thành công)", "Mỗi ứng viên có phân loại điều kiện bắt buộc; ứng viên đủ điều "
                                       "kiện có điểm từng tiêu chí, điểm tổng và thứ hạng; shortlist có ít "
                                       "nhất một ứng viên chính; tiêu chí, trọng số và kết quả của lần đánh "
                                       "giá được lưu (QT-16). Kết quả sẵn sàng cho UC-03.4 và bàn giao V5."),
        ("Hậu điều kiện (không có ứng viên đạt)",
         "Phân loại và lý do loại được lưu; không lập shortlist. Mở rộng nguồn ứng viên thuộc UC-03.2.\n"
         "[Cần bổ sung — V4: nhánh phản hồi V2 khi đã xét hết nguồn mà không có ứng viên (CN-46).]"),
    ])
    para(doc, "Luồng chính", bold=True)
    table(doc, ["Bước", "Tác nhân/Hệ thống", "Mô tả"], [
        ("1", "RM", "Chọn nhu cầu đã xác nhận và danh sách ứng viên cần đánh giá."),
        ("2", "Hệ thống", "Kiểm tra tiền điều kiện; trình bày nhu cầu, danh sách ứng viên, điều kiện bắt buộc, "
                          "tiêu chí, trọng số và thông tin năng lực của ứng viên từ hồ sơ năng lực tập trung."),
        ("3", "RM", "Đối chiếu từng ứng viên với điều kiện bắt buộc của nhu cầu."),
        ("4", "Hệ thống", "Ghi phân loại Đạt, Đạt có điều kiện hoặc Không đạt kèm lý do cho trường hợp Không "
                          "đạt; cảnh báo kỹ năng chưa xác nhận hoặc quá hạn cập nhật."),
        ("5", "RM", "Xác định tập ứng viên đủ điều kiện (Đạt, Đạt có điều kiện, hoặc được miễn trừ); ứng viên "
                    "Không đạt không được miễn trừ bị loại khỏi đánh giá."),
        ("6", "RM", "Chấm điểm từng tiêu chí cho ứng viên đủ điều kiện, có xem xét cảnh báo độ tin cậy của kỹ năng."),
        ("7", "Hệ thống", "Tính điểm tổng theo trọng số của từng tiêu chí và xếp hạng ứng viên theo điểm tổng "
                          "giảm dần; đồng hạng thì xếp theo thứ tự ưu tiên nguồn."),
        ("8", "RM", "Lập shortlist: tối thiểu một ứng viên chính, khuyến nghị thêm một đến hai ứng viên thay thế."),
        ("9", "Hệ thống", "Hiển thị mức đáp ứng và khoảng thiếu năng lực của từng ứng viên trong shortlist "
                          "(CN-42, chỉ xem)."),
        ("10", "RM", "Rà soát kết quả đánh giá và xếp hạng, xác nhận hoàn tất."),
        ("11", "Hệ thống", "Lưu tiêu chí, trọng số, điểm từng tiêu chí, điểm tổng, thứ hạng và shortlist của "
                           "lần đánh giá."),
    ], widths=[1.3, 3, 11.7])
    para(doc, "Luồng thay thế và ngoại lệ", bold=True)
    table(doc, ["Mã", "Tại bước", "Điều kiện", "Xử lý"], [
        ("A1", "3", "Ứng viên có xung đột lợi ích với khách hàng.",
         "RM loại ứng viên khỏi quá trình đánh giá và ghi lý do (CN-35, QT-14); ứng viên không vào các bước sau."),
        ("A2", "4, 5", "Ứng viên Không đạt nhưng RM đề nghị miễn trừ.",
         "RM ghi đề nghị miễn trừ kèm lý do; PMO/Quản lý danh mục phê duyệt hoặc từ chối (CN-37, QT-13). "
         "Được duyệt thì ứng viên được đánh dấu “miễn trừ” và vào chấm điểm; bị từ chối thì ứng viên bị loại."),
        ("A3", "4", "Kỹ năng ứng viên chưa xác nhận hoặc đã quá 6 tháng kể từ lần đánh giá gần nhất.",
         "Hệ thống cảnh báo; RM cân nhắc độ tin cậy khi chấm điểm (CN-36, QT-15).\n"
         "[Cần sửa — V4: dẫn CN-20 cho mốc 6 tháng.]"),
        ("A4", "7", "Hai ứng viên bằng điểm tổng.",
         "Xếp theo thứ tự ưu tiên nguồn: nhân sự nội bộ, cộng tác viên hoặc đối tác, tuyển ngoài (CN-40, QT-18)."),
        ("A5", "8", "Chỉ có một ứng viên đủ điều kiện.",
         "Shortlist chỉ có một ứng viên chính; không có ứng viên thay thế (QT-17 chỉ khuyến nghị thêm)."),
        ("E1", "1, 2", "Nhu cầu chưa ở trạng thái “đã xác nhận”.", "Hệ thống không cho đánh giá; UC kết thúc (QT-12)."),
        ("E2", "2", "Nhu cầu chưa có bộ điều kiện bắt buộc, tiêu chí, trọng số.",
         "Hệ thống yêu cầu xác lập trước ở UC-03.1; UC kết thúc."),
        ("E3", "5", "Không còn ứng viên đủ điều kiện.",
         "Không lập shortlist; lưu phân loại và lý do loại; UC kết thúc.\n"
         "[Cần bổ sung — V4: quay lại nguồn kế tiếp (UC-03.2, CN-33); hết nguồn thì phản hồi V2 (CN-46).]"),
    ], widths=[1.2, 1.6, 5, 8.2])
    para(doc, "Quy tắc nghiệp vụ và chức năng liên quan", bold=True)
    table(doc, ["ID", "Nội dung", "Vai trò trong UC"], [
        ("CN-34", "Đối chiếu điều kiện bắt buộc, phân loại, ghi lý do không đạt.", "Bước 3, 4"),
        ("CN-35", "Loại ứng viên xung đột lợi ích.", "A1"),
        ("CN-36", "Cảnh báo kỹ năng chưa xác nhận hoặc quá hạn.", "A3"),
        ("CN-37", "Đề nghị miễn trừ, PMO phê duyệt.", "A2"),
        ("CN-38", "Xem điểm từng tiêu chí và điểm tổng theo trọng số.", "Bước 6, 7"),
        ("CN-39", "Xếp hạng theo điểm phù hợp.", "Bước 7"),
        ("CN-40", "Xác định thứ tự giữa các ứng viên bằng điểm.", "A4"),
        ("CN-41", "Lập shortlist: tối thiểu một chính, khuyến nghị một đến hai thay thế.", "Bước 8"),
        ("CN-42", "Xem mức đáp ứng và khoảng thiếu năng lực (chỉ xem).", "Bước 9"),
        ("QT-12", "Chỉ nhu cầu “đã xác nhận” được đánh giá.", "Tiền điều kiện, E1"),
        ("QT-13", "Không đạt không vào shortlist chính, trừ miễn trừ có phê duyệt; miễn trừ phải có lý do và "
                  "người phê duyệt.", "A2"),
        ("QT-14", "Xung đột lợi ích bị loại.", "A1"),
        ("QT-15", "Cảnh báo độ tin cậy kỹ năng khi quá thời hạn cập nhật.\n[Cần sửa — V4: mốc 6 tháng dẫn CN-20.]",
         "A3"),
        ("QT-16", "Mỗi lần đánh giá lưu tiêu chí, trọng số, kết quả.", "Bước 11"),
        ("QT-17", "Shortlist tối thiểu một chính, khuyến nghị 1–2 thay thế.", "Bước 8, A5"),
        ("QT-18", "Thứ tự ưu tiên nguồn.", "A4"),
        ("QT-19", "Điểm phù hợp tính từ trọng số và điểm từng tiêu chí; không dùng bộ trọng số cố định.", "Bước 6, 7"),
        ("QT-20", "Khả năng tham gia và chi phí không thuộc đánh giá mức độ phù hợp.", "Ngoài phạm vi UC"),
    ], widths=[2, 10.5, 3.5])
    placeholder(doc, OWNER["UC-03.3"],
                "Sơ đồ hoạt động UC-03.3: vẽ lại bằng draw.io (AD_UC-03.3.drawio), thêm nhánh CN-46 và nút "
                "merge trước quyết định “Còn ứng viên đủ điều kiện?”.")


def qd(text):
    for n in range(1, 10):
        text = text.replace(f"chờ Q-0{n}", f"[Chờ chốt QĐ-0{n}]").replace(f"Q-0{n}", f"QĐ-0{n}")
    return text


def v10_spec(doc, s, fig_no, root):
    kv_table(doc, [
        ("Tên ca sử dụng", s["name"]), ("ID", s["id"]), ("Tác nhân chính", s["actor"]),
        ("Tác nhân phụ", s["secondary"]), ("Mục tiêu", s["goal"]), ("Sự kiện kích hoạt", s["trigger"].replace(" Không chạy một lần liên tục đến tuần kế tiếp trong AD.", "")),
        ("Tiền điều kiện", s["pre"]), ("Hậu điều kiện", s["post"]), ("Đầu vào", s["inputs"]),
        ("Đầu ra", s["outputs"]), ("Truy vết", s["cn"]),
    ])
    para(doc, "Luồng chính", bold=True)
    table(doc, ["Bước", "Tác nhân/Hệ thống", "Mô tả"], [tuple(x) for x in s["steps"]],
          widths=[1.3, 3.5, 11.2])
    para(doc, "Luồng thay thế và ngoại lệ", bold=True)
    table(doc, ["Mã", "Tại bước", "Điều kiện", "Xử lý"],
          [(a, b, c, qd(d)) for a, b, c, d in s["alternatives"]], widths=[1.2, 1.6, 5, 8.2])
    png = os.path.join(root, "drawio", "II.2_V10", f"AD_{s['id']}.png")
    figure(doc, png, f"Hình II.2-{fig_no}. Sơ đồ hoạt động {s['id']} {s['name']}.", width=15)


def build(root, out):
    cat = json.load(open(os.path.join(DATA, "catalogue.json"), encoding="utf-8"))
    fds = json.load(open(os.path.join(DATA, "fds.json"), encoding="utf-8"))
    v10 = json.load(open(os.path.join(DATA, "v10.json"), encoding="utf-8"))
    spec = importlib.util.spec_from_file_location("g3", os.path.join(HERE, "gen_ii3_drawio.py"))
    g3 = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(g3)
    ucs = {u["id"]: u for u in cat["ucs"]}
    dr = os.path.join(root, "drawio")

    doc = Document()
    setup(doc)
    doc.add_heading("II. PHÂN TÍCH (THIẾT KẾ LOGIC)", level=1)

    # ---------------- II.1
    doc.add_heading("II.1. Phạm vi hệ thống công nghệ", level=2)
    para(doc, "Hệ thống quản lý và điều phối nguồn lực dự án của Lạc Việt Consulting hỗ trợ toàn bộ quá "
              "trình từ xác định nhu cầu nguồn lực đến theo dõi và dự báo cung – cầu. Phạm vi gồm 9 chức "
              "năng cấp vùng (UC-01 đến UC-09), tương ứng các vùng nghiệp vụ V2–V10 của Phần I, và 28 chức "
              "năng con. Tác nhân gồm PM, RM, Nguồn lực, Tài chính, PMO/Quản lý danh mục và CEO/Managing "
              "Director.")
    figure(doc, os.path.join(dr, "II.1", "II1_UC_Tong_Quat.png"), "Hình II.1-1. Sơ đồ ca sử dụng tổng quát.",
           width=14)
    figure(doc, os.path.join(dr, "II.1", "Phan_cap_chuc_nang.png"),
           "Hình II.1-2. Sơ đồ phân cấp chức năng.", width=16)
    para(doc, "Bảng II.1-1 liệt kê 28 chức năng con. 12 ca sử dụng trọng tâm được đặc tả đầy đủ kèm sơ đồ "
              "hoạt động; các ca sử dụng còn lại được mô tả bằng FDS tóm tắt.")
    caption(doc, "Bảng II.1-1. Danh sách chức năng con.", before=True)
    rows = []
    for u in cat["ucs"]:
        rows.append((u["region"], u["id"], u["name"], u["actor"].replace(";", ","), "CN-" + u["cn"],
                     "Đặc tả đầy đủ, AD" if u["id"] in FOCUS else "FDS tóm tắt"))
    table(doc, ["Vùng", "Mã", "Chức năng", "Tác nhân", "Chức năng Phần I", "Mức đặc tả"], rows,
          widths=[1.3, 1.9, 5, 2.6, 3, 2.2], size=10)

    # ---------------- II.2
    doc.add_heading("II.2. Mô hình hóa chức năng của hệ thống công nghệ", level=2)
    para(doc, "Mỗi vùng gồm sơ đồ ca sử dụng chi tiết, FDS tóm tắt cho các ca sử dụng không trọng tâm, đặc "
              "tả đầy đủ và sơ đồ hoạt động cho ca sử dụng trọng tâm.")
    fig = 0
    for k, (code, vname, v, subs) in enumerate(cat["regions"], 1):
        doc.add_heading(f"II.2.{k}. {v} — {vname}", level=3)
        fig += 1
        figure(doc, os.path.join(dr, "II.2_UCD_vung", f"UCD_{code}.png"),
               f"Hình II.2-{fig}. Sơ đồ ca sử dụng {code} {vname}.", width=15)
        for cid, cname in subs:
            if cid in FOCUS:
                doc.add_heading(f"{cid} {cname} (ca sử dụng trọng tâm)", level=4)
                if cid == "UC-03.3":
                    v4_33(doc)
                elif v == "V10":
                    s = next(x for x in v10["specs"] if x["id"] == cid)
                    fig += 1
                    v10_spec(doc, s, fig, root)
                else:
                    focus_template(doc, ucs[cid])
            else:
                doc.add_heading(f"{cid} {cname} (FDS tóm tắt)", level=4)
                kv_table(doc, clean_fds(cid, fds[cid]))
        if v == "V10":
            para(doc, "Quy tắc nghiệp vụ áp dụng cho UC-09.1 và UC-09.2", bold=True)
            table(doc, ["Căn cứ", "Nội dung"], [tuple(r) for r in v10["rules"]], widths=[4.2, 11.8])

    # ---------------- II.3
    doc.add_heading("II.3. Mô hình hóa cấu trúc tĩnh của hệ thống công nghệ", level=2)
    para(doc, "Sơ đồ lớp miền nghiệp vụ gồm 48 lớp, chia thành năm gói theo nhóm vùng. Hình II.3-1 cho thấy "
              "các gói và quan hệ phụ thuộc giữa chúng. Hình II.3-2 đến II.3-6 trình bày từng gói với thuộc "
              "tính, liên kết và bội số. Lớp thuộc gói khác được vẽ nền xám, viền nét đứt, chỉ để thể hiện "
              "liên kết; thuộc tính của lớp đó xem ở hình của gói sở hữu.")
    files = ["II3-1_Tong_quan_goi_lop", "II3-2_Nhu_cau_va_nang_luc", "II3-3_Tim_kiem_kha_nang_tham_gia_lich",
             "II3-4_Phuong_an_xung_dot_phe_duyet", "II3-5_Chi_phi_ngan_sach", "II3-6_Theo_doi_du_bao"]
    figure(doc, os.path.join(dr, "II.3", files[0] + ".png"),
           "Hình II.3-1. Tổng quan các gói lớp miền nghiệp vụ.", width=15)
    for (key, fno, title, vung, _), f in zip(g3.CLUSTERS, files[1:]):
        doc.add_heading(f"Gói {key}: {title} ({vung})", level=3)
        figure(doc, os.path.join(dr, "II.3", f + ".png"), f"Hình {fno}. Sơ đồ lớp miền: {title.lower()}.",
               width=16)
    doc.add_heading("Danh mục lớp", level=3)
    caption(doc, "Bảng II.3-1. Danh mục lớp miền nghiệp vụ.", before=True)
    rows = []
    for key, fno, _, vung, cs in g3.CLUSTERS:
        for c in cs:
            attrs = ", ".join(a for a in g3.CLASSES[c])
            rows.append((c, MEANING[c], attrs, fno))
    table(doc, ["Lớp", "Ý nghĩa", "Thuộc tính", "Hình"], rows, widths=[4.4, 4.6, 5.4, 1.6], size=10)

    # ---------------- II.4
    doc.add_heading("II.4. Mô hình hóa hành vi của hệ thống công nghệ", level=2)
    para(doc, "Các sơ đồ hành vi dựa trên ca sử dụng ở II.2 và lớp miền ở II.3. Sơ đồ trạng thái mô tả vòng "
              "đời của đối tượng có hành vi phức tạp; sơ đồ giao tiếp mô tả các sự kiện truy vấn dữ liệu; sơ "
              "đồ tuần tự mô tả các sự kiện thay đổi dữ liệu.")
    doc.add_heading("II.4.1. Sơ đồ trạng thái", level=3)
    table(doc, ["Đối tượng", "Ca sử dụng", "Người thực hiện"], [(a, b, c) for a, b, c, *_ in SM],
          widths=[4.5, 5, 6.5])
    for obj, uc, who, src, f in SM:
        doc.add_heading(f"Vòng đời {obj}", level=4)
        placeholder(doc, who, f"Chèn PNG xuất từ {f}.drawio. Truy vết: {src}.")
    doc.add_heading("II.4.2. Sơ đồ giao tiếp", level=3)
    table(doc, ["Sự kiện truy vấn", "Ca sử dụng", "Người thực hiện"], [(a, b, c) for a, b, c, *_ in COM],
          widths=[6, 3, 7])
    for ev, uc, who, src, f in COM:
        doc.add_heading(f"{uc} {ev}", level=4)
        placeholder(doc, who, f"Chèn PNG xuất từ {f}.drawio. Truy vết: {src}.")
    doc.add_heading("II.4.3. Sơ đồ tuần tự", level=3)
    table(doc, ["Sự kiện thay đổi dữ liệu", "Ca sử dụng", "Người thực hiện"], [(a, b, c) for a, b, c, _ in SEQ],
          widths=[6, 3, 7])
    for ev, uc, who, src in SEQ:
        doc.add_heading(f"{uc} {ev}", level=4)
        placeholder(doc, who, f"Chèn PNG xuất từ SEQ_{uc}.drawio. Truy vết: {src}.")

    doc.save(out)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
